"""End-to-end API tests over the in-memory database.

Covers onboarding, the full check-in / check-out flow and the anti-proxy
rejections (wrong face, replayed challenge, wrong phone, weak beacon,
fake GPS, outside campus, duplicate face, …).
"""

from __future__ import annotations

import io
from typing import Any

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.conftest import (
    CAMPUS_LAT,
    CAMPUS_LNG,
    Clock,
    World,
    beacon_token,
    headers,
    near,
    vector,
)

ON_CAMPUS = {"latitude": CAMPUS_LAT + 0.0005, "longitude": CAMPUS_LNG, "accuracy_m": 15}


def reading(beacon: dict[str, Any], rssi: int = -60, windows_ago: int = 0) -> dict[str, Any]:
    return {"beacon_id": beacon["id"], "token": beacon_token(beacon, windows_ago), "rssi": rssi}


def challenge(
    client: TestClient,
    staff: dict[str, Any],
    device: dict[str, Any],
    bcn: dict[str, Any],
    action: str = "check_in",
    **overrides: Any,
) -> Any:
    body = {
        "action": action,
        "beacon": reading(bcn),
        "device_fingerprint": device["device_fingerprint"],
        **overrides,
    }
    return client.post("/api/v1/attendance/challenge", json=body, headers=headers(staff))


def submit(
    client: TestClient,
    staff: dict[str, Any],
    device: dict[str, Any],
    bcn: dict[str, Any],
    ch: dict[str, Any],
    embedding: list[float],
    **overrides: Any,
) -> Any:
    body = {
        "challenge_id": ch["challenge_id"],
        "beacon": reading(bcn),
        "embedding": embedding,
        "completed_steps": ch["liveness_steps"],
        "device_fingerprint": device["device_fingerprint"],
        "location": ON_CAMPUS,
        **overrides,
    }
    return client.post("/api/v1/attendance/submit", json=body, headers=headers(staff))


def error_code(resp: Any) -> str:
    return resp.json()["error"]["code"]


# ── Auth & onboarding ─────────────────────────────────────────────────────────


def test_requires_token(client: TestClient, world: World) -> None:
    assert client.get("/api/v1/me").status_code == 401


def test_onboarding_steps(client: TestClient, world: World) -> None:
    staff = world.staff(status="pending")
    h = headers(staff)
    fp = "f" * 64

    me = client.get("/api/v1/me", headers=h).json()
    assert me["onboarding"]["next_step"] == "register_device"

    r = client.post(
        "/api/v1/devices/register",
        json={"device_fingerprint": fp, "device_model": "Pixel 8"},
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert client.get("/api/v1/me", headers=h).json()["onboarding"]["next_step"] == "enroll_face"

    base = vector(11)
    r = client.post(
        "/api/v1/face/enroll",
        json={
            "embeddings": [near(base, i) for i in range(3)],
            "model_version": "mobilefacenet-v1",
            "device_fingerprint": fp,
        },
        headers=h,
    )
    assert r.status_code == 200, r.text
    assert r.json()["status"] == "pending"
    assert "embeddings" not in r.json()
    me = client.get("/api/v1/me", headers=h).json()
    assert me["onboarding"]["next_step"] == "await_approval"

    # Pending staff cannot mark attendance.
    beacon = world.beacon()
    r = challenge(client, staff, {"device_fingerprint": fp}, beacon)
    assert r.status_code == 403 and error_code(r) == "ACCOUNT_PENDING"

    admin = world.staff("Admin One", role="admin")
    r = client.post(f"/api/v1/admin/staff/{staff['id']}/approve", headers=headers(admin))
    assert r.status_code == 200, r.text
    me = client.get("/api/v1/me", headers=h).json()
    assert me["onboarding"]["next_step"] == "ready"
    assert me["face"]["status"] == "approved"


def test_device_binding_rules(client: TestClient, world: World) -> None:
    a = world.staff("Asha Rao")
    b = world.staff("Bala Kumar")
    fp_a, fp_b = "a" * 64, "b" * 64
    reg = "/api/v1/devices/register"

    assert (
        client.post(reg, json={"device_fingerprint": fp_a}, headers=headers(a)).status_code == 200
    )
    # Same phone again → fine (metadata refresh).
    assert (
        client.post(reg, json={"device_fingerprint": fp_a}, headers=headers(a)).status_code == 200
    )
    # A second phone for the same account → refused.
    r = client.post(reg, json={"device_fingerprint": fp_b}, headers=headers(a))
    assert r.status_code == 409 and error_code(r) == "DEVICE_ALREADY_BOUND"
    # Someone else's phone → refused.
    r = client.post(reg, json={"device_fingerprint": fp_a}, headers=headers(b))
    assert r.status_code == 409 and error_code(r) == "DEVICE_IN_USE"
    # Emulator → refused.
    r = client.post(
        reg, json={"device_fingerprint": fp_b, "is_physical_device": False}, headers=headers(b)
    )
    assert r.status_code == 422 and error_code(r) == "EMULATOR_NOT_ALLOWED"

    # Admin reset lets A bind a new phone.
    admin = world.staff("Admin One", role="admin")
    r = client.post(f"/api/v1/admin/staff/{a['id']}/reset-device", headers=headers(admin))
    assert r.json() == {"devices_unbound": 1}
    assert (
        client.post(reg, json={"device_fingerprint": fp_b}, headers=headers(a)).status_code == 200
    )


def test_enroll_rejects_inconsistent_and_duplicate_faces(client: TestClient, world: World) -> None:
    # Bala is already enrolled & approved.
    world.ready_staff(seed=21, name="Bala Kumar")
    asha = world.staff("Asha Rao", status="pending")
    device = world.device(asha["id"])
    url = "/api/v1/face/enroll"
    body = {"model_version": "mobilefacenet-v1", "device_fingerprint": device["device_fingerprint"]}

    # Three different people → inconsistent.
    r = client.post(
        url, json={**body, "embeddings": [vector(1), vector(2), vector(3)]}, headers=headers(asha)
    )
    assert r.status_code == 422 and error_code(r) == "FACE_INCONSISTENT"

    # Bala's face on Asha's account → duplicate (proxy enrollment).
    bala = vector(21)
    r = client.post(
        url,
        json={**body, "embeddings": [near(bala, 100 + i) for i in range(3)]},
        headers=headers(asha),
    )
    assert r.status_code == 409 and error_code(r) == "FACE_DUPLICATE"

    # Too few samples.
    r = client.post(url, json={**body, "embeddings": [vector(5)]}, headers=headers(asha))
    assert r.status_code == 422 and error_code(r) == "FACE_SAMPLES"


# ── Check-in / check-out ──────────────────────────────────────────────────────


def test_full_check_in_and_check_out(client: TestClient, world: World, clock: Clock) -> None:
    clock.set_local(9, 5)
    staff, device, base = world.ready_staff()
    beacon = world.beacon()

    r = challenge(client, staff, device, beacon)
    assert r.status_code == 200, r.text
    ch = r.json()
    assert len(ch["liveness_steps"]) == 2 and ch["beacon"]["name"] == "Staff Room"

    r = submit(client, staff, device, beacon, ch, near(base, 500))
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["attendance"]["status"] == "present"
    assert body["verification"]["face_score"] > 0.6
    assert body["verification"]["distance_m"] < 100

    today = client.get("/api/v1/attendance/today", headers=headers(staff)).json()
    assert today["attendance"]["check_in_at"]

    # Second check-in is refused at the challenge step.
    r = challenge(client, staff, device, beacon)
    assert r.status_code == 422 and error_code(r) == "ALREADY_CHECKED_IN"

    clock.set_local(17, 30)
    ch = challenge(client, staff, device, beacon, action="check_out").json()
    r = submit(client, staff, device, beacon, ch, near(base, 501))
    assert r.status_code == 200, r.text
    assert r.json()["attendance"]["check_out_at"]

    r = challenge(client, staff, device, beacon, action="check_out")
    assert error_code(r) == "ALREADY_CHECKED_OUT"

    history = client.get("/api/v1/attendance/history", headers=headers(staff)).json()
    assert len(history) == 1

    attempts = world.db.rows("attendance_attempts")
    assert sum(1 for a in attempts if a["success"]) == 2


def test_late_after_grace(client: TestClient, world: World, clock: Clock) -> None:
    clock.set_local(9, 31)
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    r = submit(client, staff, device, beacon, ch, near(base, 500))
    assert r.json()["attendance"]["status"] == "late"


def test_check_out_requires_check_in(client: TestClient, world: World) -> None:
    staff, device, _ = world.ready_staff()
    r = challenge(client, staff, device, world.beacon(), action="check_out")
    assert error_code(r) == "NOT_CHECKED_IN"


# ── Anti-proxy rejections ─────────────────────────────────────────────────────


def test_wrong_face_is_rejected_and_logged(client: TestClient, world: World) -> None:
    staff, device, _ = world.ready_staff(seed=7)
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    r = submit(client, staff, device, beacon, ch, vector(999))
    assert r.status_code == 422 and error_code(r) == "FACE_MISMATCH"
    failed = [a for a in world.db.rows("attendance_attempts") if not a["success"]]
    assert failed[-1]["reason_code"] == "FACE_MISMATCH"
    assert failed[-1]["face_score"] < 0.6
    assert world.db.rows("attendance") == []


def test_challenge_cannot_be_replayed(client: TestClient, world: World) -> None:
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    # First submit fails (wrong face) — the challenge is burned anyway.
    assert error_code(submit(client, staff, device, beacon, ch, vector(999))) == "FACE_MISMATCH"
    r = submit(client, staff, device, beacon, ch, near(base, 1))
    assert error_code(r) == "CHALLENGE_USED"


def test_challenge_belongs_to_one_staff(client: TestClient, world: World) -> None:
    a, dev_a, _ = world.ready_staff(seed=7, name="Asha Rao")
    b, dev_b, base_b = world.ready_staff(seed=8, name="Bala Kumar")
    beacon = world.beacon()
    ch = challenge(client, a, dev_a, beacon).json()
    r = submit(client, b, dev_b, beacon, ch, near(base_b, 1))
    assert error_code(r) == "CHALLENGE_INVALID"


def test_expired_challenge(client: TestClient, world: World, clock: Clock) -> None:
    clock.set_local(9, 0)
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    clock.set_local(9, 5)
    assert error_code(submit(client, staff, device, beacon, ch, near(base, 1))) == (
        "CHALLENGE_EXPIRED"
    )


def test_wrong_liveness_steps(client: TestClient, world: World) -> None:
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    r = submit(
        client,
        staff,
        device,
        beacon,
        ch,
        near(base, 1),
        completed_steps=list(reversed(ch["liveness_steps"])),
    )
    assert error_code(r) == "LIVENESS_FAILED"


def test_other_phone_is_rejected(client: TestClient, world: World) -> None:
    staff, _device, _ = world.ready_staff()
    r = challenge(client, staff, {"device_fingerprint": "z" * 64}, world.beacon())
    assert error_code(r) == "DEVICE_MISMATCH"


def test_beacon_checks(client: TestClient, world: World) -> None:
    staff, device, _ = world.ready_staff()
    beacon = world.beacon(rssi_threshold=-80)

    weak = challenge(client, staff, device, beacon, beacon=reading(beacon, rssi=-92))
    assert error_code(weak) == "BEACON_TOO_FAR"

    stale = challenge(client, staff, device, beacon, beacon=reading(beacon, windows_ago=4))
    assert error_code(stale) == "BEACON_TOKEN_INVALID"

    forged = {**reading(beacon), "token": "0123456789abcdef"}
    assert error_code(challenge(client, staff, device, beacon, beacon=forged)) == (
        "BEACON_TOKEN_INVALID"
    )

    unknown = {**reading(beacon), "beacon_id": "11111111-2222-3333-4444-555555555555"}
    assert error_code(challenge(client, staff, device, beacon, beacon=unknown)) == "BEACON_UNKNOWN"


def test_mock_location_rejected(client: TestClient, world: World) -> None:
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    r = submit(
        client, staff, device, beacon, ch, near(base, 1), location={**ON_CAMPUS, "is_mocked": True}
    )
    assert error_code(r) == "MOCK_LOCATION"


def test_geofence_enforce_vs_flag(client: TestClient, world: World) -> None:
    staff, device, base = world.ready_staff()
    beacon = world.beacon()
    far = {"latitude": CAMPUS_LAT + 0.2, "longitude": CAMPUS_LNG, "accuracy_m": 10}

    world.db.rows("campus_settings")[0]["geofence_mode"] = "enforce"
    ch = challenge(client, staff, device, beacon).json()
    r = submit(client, staff, device, beacon, ch, near(base, 1), location=far)
    assert error_code(r) == "OUTSIDE_CAMPUS"

    world.db.rows("campus_settings")[0]["geofence_mode"] = "flag"
    ch = challenge(client, staff, device, beacon).json()
    r = submit(client, staff, device, beacon, ch, near(base, 1), location=far)
    assert r.status_code == 200
    assert "outside_geofence" in r.json()["attendance"]["flags"]


def test_emulator_rejected(client: TestClient, world: World) -> None:
    staff, device, _ = world.ready_staff()
    r = challenge(client, staff, device, world.beacon(), is_physical_device=False)
    assert error_code(r) == "EMULATOR_NOT_ALLOWED"


def test_unapproved_face_cannot_mark(client: TestClient, world: World) -> None:
    staff = world.staff()
    device = world.device(staff["id"])
    world.face(staff["id"], 3, status="pending")
    r = challenge(client, staff, device, world.beacon())
    assert error_code(r) == "FACE_NOT_APPROVED"


# ── Admin ─────────────────────────────────────────────────────────────────────


def test_admin_only(client: TestClient, world: World) -> None:
    staff = world.staff()
    assert client.get("/api/v1/admin/staff", headers=headers(staff)).status_code == 403


def test_admin_approve_blocks_duplicate_face(client: TestClient, world: World) -> None:
    admin = world.staff("Admin One", role="admin")
    world.ready_staff(seed=31, name="Bala Kumar")
    proxy = world.staff("Proxy Person", status="pending")
    world.device(proxy["id"])
    # Template copied from Bala's face, still pending.
    world.face(proxy["id"], 31, status="pending")
    r = client.post(f"/api/v1/admin/staff/{proxy['id']}/approve", headers=headers(admin))
    assert r.status_code == 409 and error_code(r) == "FACE_DUPLICATE"
    assert "Bala Kumar" in r.json()["error"]["message"]


def test_roster_manual_and_export(client: TestClient, world: World, clock: Clock) -> None:
    clock.set_local(9, 0)
    admin = world.staff("Admin One", role="admin")
    staff, device, base = world.ready_staff(name="Asha Rao")
    absent = world.staff("Chitra Devi")
    beacon = world.beacon()
    ch = challenge(client, staff, device, beacon).json()
    assert submit(client, staff, device, beacon, ch, near(base, 1)).status_code == 200

    roster = client.get("/api/v1/admin/attendance", headers=headers(admin)).json()
    assert roster["summary"]["present"] == 1
    assert roster["summary"]["not_marked"] == 2  # admin + Chitra

    r = client.post(
        "/api/v1/admin/attendance/manual",
        json={
            "staff_id": absent["id"],
            "date": roster["date"],
            "status": "on_leave",
            "reason": "Medical leave",
        },
        headers=headers(admin),
    )
    assert r.status_code == 200, r.text
    assert r.json()["is_manual"] is True

    # On-leave staff cannot check in that day.
    world.device(absent["id"])
    world.face(absent["id"], 77)
    dev = world.db.rows("devices")[-1]
    assert error_code(challenge(client, absent, dev, beacon)) == "DAY_LOCKED"

    r = client.get(
        f"/api/v1/admin/reports/export?from={roster['date']}&to={roster['date']}",
        headers=headers(admin),
    )
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Daily log", "Summary", "Info"]
    statuses = {row[2]: row[4] for row in wb["Daily log"].iter_rows(min_row=2, values_only=True)}
    assert statuses["Asha Rao"] == "Present"
    assert statuses["Chitra Devi"] == "On leave"

    audit = client.get("/api/v1/admin/audit", headers=headers(admin)).json()
    assert {a["action"] for a in audit} >= {"attendance_manual", "report_export"}


def test_attempts_feed(client: TestClient, world: World) -> None:
    admin = world.staff("Admin One", role="admin")
    staff, _device, _ = world.ready_staff(name="Asha Rao")
    challenge(client, staff, {"device_fingerprint": "z" * 64}, world.beacon())
    rows = client.get("/api/v1/admin/attempts", headers=headers(admin)).json()
    assert rows and rows[0]["reason_code"] == "DEVICE_MISMATCH"
    assert rows[0]["staff_name"] == "Asha Rao"
    assert "device_fingerprint" not in rows[0]


def test_beacon_admin_flow(client: TestClient, world: World) -> None:
    admin = world.staff("Admin One", role="admin")
    h = headers(admin)
    r = client.post("/api/v1/admin/beacons", json={"name": "Staff Room A"}, headers=h)
    assert r.status_code == 201
    created = r.json()
    assert len(created["secret"]) == 64
    assert created["id"] in created["firmware_config"]
    listed = client.get("/api/v1/admin/beacons", headers=h).json()
    assert "secret" not in listed[0]
    rotated = client.post(f"/api/v1/admin/beacons/{created['id']}/rotate-secret", headers=h).json()
    assert rotated["secret"] != created["secret"]


def test_campus_settings_update(client: TestClient, world: World) -> None:
    admin = world.staff("Admin One", role="admin")
    r = client.put(
        "/api/v1/admin/campus",
        json={"geofence_mode": "enforce", "work_start_time": "08:30", "radius_m": 800},
        headers=headers(admin),
    )
    assert r.status_code == 200, r.text
    assert r.json()["geofence_mode"] == "enforce"
    assert r.json()["work_start_time"].startswith("08:30")
    bad = client.put("/api/v1/admin/campus", json={"timezone": "Mars/Base"}, headers=headers(admin))
    assert error_code(bad) == "INVALID_TIMEZONE"


def test_admin_cannot_lock_self_out(client: TestClient, world: World) -> None:
    admin = world.staff("Admin One", role="admin")
    r = client.post(
        f"/api/v1/admin/staff/{admin['id']}/role", json={"role": "staff"}, headers=headers(admin)
    )
    assert error_code(r) == "SELF_ACTION"


def test_validation_envelope(client: TestClient, world: World) -> None:
    staff = world.staff()
    r = client.post(
        "/api/v1/attendance/challenge", json={"action": "teleport"}, headers=headers(staff)
    )
    assert r.status_code == 422 and error_code(r) == "VALIDATION_ERROR"


def test_health(client: TestClient) -> None:
    assert client.get("/health").json()["status"] == "healthy"
