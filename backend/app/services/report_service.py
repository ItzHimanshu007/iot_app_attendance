"""Excel export of staff attendance (openpyxl)."""

from __future__ import annotations

import io
from datetime import date, timedelta
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.core.exceptions import ValidationError
from app.core.timeutil import campus_zone, parse_ts
from app.repositories.attendance_repo import AttendanceRepository
from app.repositories.staff_repo import StaffRepository
from app.services import campus_service

MAX_DAYS = 93

_attendance = AttendanceRepository()
_staff = StaffRepository()

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4F46E5")
_STATUS_LABEL = {
    "present": "Present",
    "late": "Late",
    "absent": "Absent",
    "on_leave": "On leave",
    "not_marked": "Not marked",
}


def _fmt_time(value: Any, tz: Any) -> str:
    ts = parse_ts(value)
    return ts.astimezone(tz).strftime("%I:%M %p") if ts else ""


def _hours(check_in: Any, check_out: Any) -> float | str:
    a, b = parse_ts(check_in), parse_ts(check_out)
    if not a or not b:
        return ""
    return round((b - a).total_seconds() / 3600, 2)


def _style_header(ws: Any, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws: Any) -> None:
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 45)


def build_export(date_from: date, date_to: date) -> bytes:
    """Workbook with a day-by-day log and a per-staff summary."""
    if date_from > date_to:
        raise ValidationError("'from' must be before 'to'", "INVALID_RANGE")
    days = (date_to - date_from).days + 1
    if days > MAX_DAYS:
        raise ValidationError(f"Export at most {MAX_DAYS} days at a time", "RANGE_TOO_LARGE")

    campus = campus_service.get_campus()
    tz = campus_zone(campus.get("timezone"))
    records = {
        (r["staff_id"], r["attendance_date"]): r
        for r in _attendance.list_range(date_from.isoformat(), date_to.isoformat())
    }
    staff_with_records = {sid for sid, _ in records}
    staff = [
        s
        for s in _staff.list_staff(limit=5000)
        if s["status"] == "active" or s["id"] in staff_with_records
    ]

    wb = Workbook()
    log = wb.active
    log.title = "Daily log"
    _style_header(
        log,
        [
            "Date",
            "Employee ID",
            "Name",
            "Department",
            "Status",
            "Check-in",
            "Check-out",
            "Hours",
            "Face score",
            "Flags",
            "Manual",
            "Reason",
        ],
    )

    totals: dict[str, dict[str, int]] = {s["id"]: dict.fromkeys(_STATUS_LABEL, 0) for s in staff}
    for offset in range(days):
        day = (date_from + timedelta(days=offset)).isoformat()
        for s in staff:
            r = records.get((s["id"], day))
            status = r["status"] if r else "not_marked"
            totals[s["id"]][status] += 1
            log.append(
                [
                    day,
                    s.get("employee_id") or "",
                    s.get("full_name"),
                    s.get("department") or "",
                    _STATUS_LABEL[status],
                    _fmt_time(r.get("check_in_at"), tz) if r else "",
                    _fmt_time(r.get("check_out_at"), tz) if r else "",
                    _hours(r.get("check_in_at"), r.get("check_out_at")) if r else "",
                    r.get("check_in_face_score") if r else "",
                    ", ".join(r.get("flags") or []) if r else "",
                    "Yes" if r and r.get("is_manual") else "",
                    (r.get("manual_reason") or "") if r else "",
                ]
            )
    _autosize(log)

    summary = wb.create_sheet("Summary")
    _style_header(
        summary,
        ["Employee ID", "Name", "Department", *[_STATUS_LABEL[k] for k in _STATUS_LABEL]],
    )
    for s in staff:
        summary.append(
            [
                s.get("employee_id") or "",
                s.get("full_name"),
                s.get("department") or "",
                *[totals[s["id"]][k] for k in _STATUS_LABEL],
            ]
        )
    _autosize(summary)

    info = wb.create_sheet("Info")
    info.append(["Campus", campus.get("campus_name")])
    info.append(["From", date_from.isoformat()])
    info.append(["To", date_to.isoformat()])
    info.append(["Timezone", str(tz)])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
