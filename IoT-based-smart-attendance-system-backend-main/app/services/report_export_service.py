"""Report export service — generates Excel files for completed session reports."""

from __future__ import annotations

from datetime import datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font


def format_datetime(iso_str: str | None) -> str:
    """Format an ISO timestamp to DD-MMM-YYYY hh:mm AM/PM format."""
    if not iso_str:
        return ""
    try:
        cleaned = iso_str.replace("Z", "+00:00")
        dt = datetime.fromisoformat(cleaned)
        return dt.strftime("%d-%b-%Y %I:%M %p")
    except Exception:
        return iso_str


def build_session_report_workbook(report_data: dict[str, Any]) -> Workbook:
    """Generate openpyxl Workbook representing a completed session compliance report.

    Performs formatting only. No attendance calculations are made here.
    """
    wb = Workbook()

    # Configure workbook properties
    wb.properties.title = "Session Report"
    wb.properties.creator = "Smart Campus Attendance System"
    wb.properties.subject = "Attendance Report"
    wb.properties.description = "Generated attendance report for completed session."

    # Use the active sheet and rename it
    ws = wb.active
    ws.title = "Session Report"

    # Setup styles
    bold_font = Font(bold=True)
    center_align = Alignment(horizontal="center")

    session = report_data.get("session") or {}
    summary = report_data.get("summary") or {}
    students = report_data.get("students") or []

    # 1. Populate metadata rows (1-5)
    meta_rows = [
        ("Subject", session.get("subject_name", "Unknown Subject")),
        ("Teacher", session.get("teacher_name", "Unknown Teacher")),
        ("Classroom", session.get("classroom_name", "Unknown Classroom")),
        ("Started", format_datetime(session.get("started_at"))),
        ("Completed", format_datetime(session.get("completed_at"))),
    ]

    for idx, (label, val) in enumerate(meta_rows, start=1):
        cell_lbl = ws.cell(row=idx, column=1, value=label)
        cell_lbl.font = bold_font
        ws.cell(row=idx, column=2, value=val)

    # Blank row 6 is empty by default

    # 2. Populate Summary rows (7-12)
    cell_sum = ws.cell(row=7, column=1, value="Summary")
    cell_sum.font = bold_font

    summary_rows = [
        ("Present", summary.get("present_count", 0)),
        ("Late", summary.get("late_count", 0)),
        ("Absent", summary.get("absent_count", 0)),
        ("Manual", summary.get("manual_count", 0)),
        ("Total Students", summary.get("total_students", 0)),
    ]

    for idx, (label, val) in enumerate(summary_rows, start=8):
        cell_lbl = ws.cell(row=idx, column=1, value=label)
        cell_lbl.font = bold_font
        ws.cell(row=idx, column=2, value=val)

    # Blank row 13 is empty by default

    # 3. Populate Spreadsheet headers (row 14)
    headers = [
        "Roll Number",
        "Student Name",
        "Attendance Status",
        "Verification Method",
        "Time Marked",
    ]

    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=14, column=col_idx, value=header)
        cell.font = bold_font
        # Center-align all headers except Student Name (column 2)
        if col_idx != 2:
            cell.alignment = center_align

    # 4. Populate Student rows (row 15 onwards)
    for row_idx, student in enumerate(students, start=15):
        ws.cell(row=row_idx, column=1, value=student.get("roll_number"))
        ws.cell(row=row_idx, column=2, value=student.get("student_name"))
        ws.cell(row=row_idx, column=3, value=student.get("attendance_status"))
        ws.cell(row=row_idx, column=4, value=student.get("verification_method"))

        marked_at = student.get("marked_at")
        ws.cell(row=row_idx, column=5, value=format_datetime(marked_at))

        # Center-align columns 1, 3, 4, 5
        for col_idx in (1, 3, 4, 5):
            ws.cell(row=row_idx, column=col_idx).alignment = center_align

    # 5. Apply Scroll Freezing below header row
    ws.freeze_panes = "A15"

    # 6. Apply column auto-sizing
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    return wb
