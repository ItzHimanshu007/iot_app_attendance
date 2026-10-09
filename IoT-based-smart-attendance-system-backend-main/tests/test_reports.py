from __future__ import annotations

import io
from unittest.mock import patch

import openpyxl


def _make_mock_data():
    """Helper to generate standard mocks conforming to test requirements."""
    students = [
        {
            "id": "student-1",
            "student_id_number": "22BCS001",
            "full_name": "Rahul Sharma",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-2",
            "student_id_number": "22BCS002",
            "full_name": "Priya Patel",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-3",
            "student_id_number": "22BCS003",
            "full_name": "Amit Singh",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-4",
            "student_id_number": "22BCS004",
            "full_name": "Neha Gupta",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-5",
            "student_id_number": "22BCS005",
            "full_name": "Zero Student",
            "role": "student",
            "is_active": True,
        },
    ]

    enrollments = [
        {"student_id": "student-1", "subject_id": "subj-1", "status": "active"},
        {"student_id": "student-2", "subject_id": "subj-2", "status": "active"},
        {"student_id": "student-3", "subject_id": "subj-3", "status": "active"},
        {"student_id": "student-4", "subject_id": "subj-4", "status": "active"},
        {"student_id": "student-5", "subject_id": "subj-5", "status": "active"},
    ]

    sessions = []
    records = []

    # Student 1: 10 sessions expected, 10 present -> 100% -> normal
    for i in range(10):
        sess_id = f"sess-1-{i}"
        sessions.append({"id": sess_id, "subject_id": "subj-1", "status": "completed"})
        records.append({"session_id": sess_id, "student_id": "student-1", "status": "present"})

    # Student 2: 10 sessions expected, 8 present -> 80% -> normal
    for i in range(10):
        sess_id = f"sess-2-{i}"
        sessions.append({"id": sess_id, "subject_id": "subj-2", "status": "completed"})
        if i < 8:
            records.append({"session_id": sess_id, "student_id": "student-2", "status": "present"})
        else:
            records.append({"session_id": sess_id, "student_id": "student-2", "status": "absent"})

    # Student 3: 50 sessions expected, 37 late/present -> 74% -> warning
    for i in range(50):
        sess_id = f"sess-3-{i}"
        sessions.append({"id": sess_id, "subject_id": "subj-3", "status": "completed"})
        if i < 37:
            records.append({"session_id": sess_id, "student_id": "student-3", "status": "late"})
        else:
            records.append({"session_id": sess_id, "student_id": "student-3", "status": "absent"})

    # Student 4: 100 sessions expected, 49 present -> 49% -> critical
    for i in range(100):
        sess_id = f"sess-4-{i}"
        sessions.append({"id": sess_id, "subject_id": "subj-4", "status": "completed"})
        if i < 49:
            records.append({"session_id": sess_id, "student_id": "student-4", "status": "present"})
        else:
            records.append({"session_id": sess_id, "student_id": "student-4", "status": "revoked"})

    return students, enrollments, sessions, records


def _make_phase2_mock_data():
    """Helper to generate standard Phase 2 mocks conforming to test requirements."""
    students = [
        {
            "id": "student-1",
            "student_id_number": "22BCS001",
            "full_name": "Rahul Sharma",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-2",
            "student_id_number": "22BCS002",
            "full_name": "Priya Patel",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-3",
            "student_id_number": "22BCS003",
            "full_name": "Amit Singh",
            "role": "student",
            "is_active": True,
        },
    ]

    enrollments = [
        {"student_id": "student-1", "subject_id": "subj-1", "status": "active"},
        {"student_id": "student-2", "subject_id": "subj-1", "status": "active"},
    ]

    sessions = [
        {
            "id": "sess-sub-1",
            "subject_id": "subj-1",
            "classroom_id": "class-101",
            "teacher_id": "teacher-smith",
            "started_at": "2026-07-12T09:00:00Z",
            "ended_at": "2026-07-12T10:00:00Z",
            "status": "completed",
            "classroom": {"name": "Lab 201"},
            "teacher": {"full_name": "Dr. Smith"},
            "subject": {"name": "Data Structures"},
        },
        {
            "id": "sess-cust-2",
            "subject_id": None,
            "classroom_id": "class-102",
            "teacher_id": "teacher-jones",
            "started_at": "2026-07-12T11:00:00Z",
            "ended_at": "2026-07-12T12:00:00Z",
            "status": "completed",
            "classroom": {"name": "Seminar Room"},
            "teacher": {"full_name": "Prof. Jones"},
            "subject": None,
        },
    ]

    records = [
        {
            "session_id": "sess-sub-1",
            "student_id": "student-1",
            "status": "present",
            "verification_method": "automatic",
            "marked_at": "2026-07-12T09:05:00Z",
        },
        {
            "session_id": "sess-cust-2",
            "student_id": "student-1",
            "status": "present",
            "verification_method": "manual",
            "marked_at": "2026-07-12T11:15:00Z",
        },
        {
            "session_id": "sess-cust-2",
            "student_id": "student-2",
            "status": "late",
            "verification_method": "automatic",
            "marked_at": "2026-07-12T11:25:00Z",
        },
    ]

    return students, enrollments, sessions, records


# Shortcuts for mock targets
USER_REPO_MOCK = "app.services.report_service._user_repo.get_by_role"
SESS_REPO_MOCK = "app.services.report_service._session_repo.list_completed"
ENRL_REPO_MOCK = "app.services.report_service._enrollment_repo.list_active"
ATTD_REPO_MOCK = "app.services.report_service._attendance_repo.list_by_sessions"
SESS_REL_MOCK = "app.services.report_service._session_repo.get_completed_with_relations"


# ── Authorization Tests ────────────────────────────────────────────────────────


def test_reports_teacher_access_succeeds(teacher_client):
    """Teachers should receive HTTP 200 OK when requesting reports."""
    students, enrollments, sessions, records = _make_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/students")
        assert response.status_code == 200


def test_reports_student_access_forbidden(student_client):
    """Students should receive HTTP 403 Forbidden when requesting reports."""
    response = student_client.get("/api/v1/reports/students")
    assert response.status_code == 403
    assert response.json()["error"]["code"] == "FORBIDDEN"


# ── Calculation and Classification Tests ───────────────────────────────────────


def test_reports_student_compliance_categories(teacher_client):
    """Verify that reports correctly calculate percentages and map risk categories."""
    students, enrollments, sessions, records = _make_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/students")
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 5

        # Create lookup mapping for assertions
        by_id = {item["student_id"]: item for item in data}

        # Student 1: 100% -> normal
        s1 = by_id["student-1"]
        assert s1["student_name"] == "Rahul Sharma"
        assert s1["roll_number"] == "22BCS001"
        assert s1["total_sessions"] == 10
        assert s1["sessions_attended"] == 10
        assert s1["sessions_missed"] == 0
        assert s1["attendance_rate"] == 1.0
        assert s1["category"] == "normal"

        # Student 2: 80% -> normal
        s2 = by_id["student-2"]
        assert s2["student_name"] == "Priya Patel"
        assert s2["roll_number"] == "22BCS002"
        assert s2["total_sessions"] == 10
        assert s2["sessions_attended"] == 8
        assert s2["sessions_missed"] == 2
        assert s2["attendance_rate"] == 0.8
        assert s2["category"] == "normal"

        # Student 3: 74% -> warning
        s3 = by_id["student-3"]
        assert s3["student_name"] == "Amit Singh"
        assert s3["roll_number"] == "22BCS003"
        assert s3["total_sessions"] == 50
        assert s3["sessions_attended"] == 37
        assert s3["sessions_missed"] == 13
        assert s3["attendance_rate"] == 0.74
        assert s3["category"] == "warning"

        # Student 4: 49% -> critical
        s4 = by_id["student-4"]
        assert s4["student_name"] == "Neha Gupta"
        assert s4["roll_number"] == "22BCS004"
        assert s4["total_sessions"] == 100
        assert s4["sessions_attended"] == 49
        assert s4["sessions_missed"] == 51
        assert s4["attendance_rate"] == 0.49
        assert s4["category"] == "critical"

        # Student 5: Zero completed sessions -> category: normal
        s5 = by_id["student-5"]
        assert s5["student_name"] == "Zero Student"
        assert s5["roll_number"] == "22BCS005"
        assert s5["total_sessions"] == 0
        assert s5["sessions_attended"] == 0
        assert s5["sessions_missed"] == 0
        assert s5["attendance_rate"] == 0.0
        assert s5["category"] == "normal"


# ── Search & Filter Tests ──────────────────────────────────────────────────────


def test_reports_search_by_name(teacher_client):
    """Verify that report generation filters matching names (case-insensitive)."""
    students, enrollments, sessions, records = _make_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        # Searching "rahul" (case-insensitive) should match Rahul Sharma
        response = teacher_client.get("/api/v1/reports/students", params={"search": "rahul"})
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 1
        assert data[0]["student_id"] == "student-1"


def test_reports_search_by_roll_number(teacher_client):
    """Verify that report generation filters matching roll numbers (case-insensitive)."""
    students, enrollments, sessions, records = _make_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        # Searching "22bcs003" should match Amit Singh
        response = teacher_client.get("/api/v1/reports/students", params={"search": "22bcs003"})
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 1
        assert data[0]["student_id"] == "student-3"


def test_reports_category_filter(teacher_client):
    """Verify filtering by compliance category."""
    students, enrollments, sessions, records = _make_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        # Test normal category -> returns student 1, 2, and 5
        response = teacher_client.get("/api/v1/reports/students", params={"category": "normal"})
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 3
        ids = {item["student_id"] for item in data}
        assert ids == {"student-1", "student-2", "student-5"}

        # Test warning category -> returns student 3
        response = teacher_client.get("/api/v1/reports/warning", params={"category": "warning"})
        # Note: the test routes to /api/v1/reports/students in FastAPI
        response = teacher_client.get("/api/v1/reports/students", params={"category": "warning"})
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 1
        assert data[0]["student_id"] == "student-3"

        # Test critical category -> returns student 4
        response = teacher_client.get("/api/v1/reports/students", params={"category": "critical"})
        assert response.status_code == 200
        data = response.json()
        assert len(data) == 1
        assert data[0]["student_id"] == "student-4"


# ── Custom Sessions Tests ──────────────────────────────────────────────────────


def test_reports_custom_sessions_handling(teacher_client):
    """Verify custom sessions (no subject) do not use virtual rosters or default absences."""
    students = [
        {
            "id": "student-1",
            "student_id_number": "22BCS001",
            "full_name": "Rahul Sharma",
            "role": "student",
            "is_active": True,
        },
        {
            "id": "student-2",
            "student_id_number": "22BCS002",
            "full_name": "Priya Patel",
            "role": "student",
            "is_active": True,
        },
    ]

    enrollments = []  # No enrollment mapping since it's a custom session

    # Custom session: subject_id is null/None
    sessions = [{"id": "sess-custom-1", "subject_id": None, "status": "completed"}]

    # Only Student 1 attended the custom session
    records = [{"session_id": "sess-custom-1", "student_id": "student-1", "status": "present"}]

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/students")
        assert response.status_code == 200
        data = response.json()

        by_id = {item["student_id"]: item for item in data}

        # Student 1 attended custom session -> expected 1 session
        s1 = by_id["student-1"]
        assert s1["total_sessions"] == 1
        assert s1["sessions_attended"] == 1
        assert s1["sessions_missed"] == 0
        assert s1["attendance_rate"] == 1.0

        # Student 2 did not attend -> zero session student -> category: normal
        s2 = by_id["student-2"]
        assert s2["total_sessions"] == 0
        assert s2["sessions_attended"] == 0
        assert s2["sessions_missed"] == 0
        assert s2["attendance_rate"] == 0.0
        assert s2["category"] == "normal"


# ── Phase 2: Session Reports API Tests ──────────────────────────────────────────


def test_reports_sessions_list_succeeds(teacher_client):
    """Verify listing completed sessions returns correct paginated statistics."""
    students, enrollments, sessions, records = _make_phase2_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get(
            "/api/v1/reports/sessions", params={"page": 1, "page_size": 10}
        )
        assert response.status_code == 200
        data = response.json()
        assert "items" in data
        assert data["total"] == 2
        assert len(data["items"]) == 2

        # Check sorting: completed_at DESC -> sess-cust-2 (completed at 12:00:00) comes first
        s2 = data["items"][0]
        assert s2["session_id"] == "sess-cust-2"
        assert s2["subject_name"] == "Unknown Subject"
        assert s2["classroom_name"] == "Seminar Room"
        assert s2["teacher_name"] == "Prof. Jones"
        assert s2["present_count"] == 1
        assert s2["late_count"] == 1
        assert s2["absent_count"] == 0
        assert s2["manual_count"] == 1
        assert s2["total_students"] == 2

        # completed session 1 comes second
        s1 = data["items"][1]
        assert s1["session_id"] == "sess-sub-1"
        assert s1["subject_name"] == "Data Structures"
        assert s1["classroom_name"] == "Lab 201"
        assert s1["teacher_name"] == "Dr. Smith"
        assert s1["present_count"] == 1
        assert s1["late_count"] == 0
        assert s1["absent_count"] == 1
        assert s1["manual_count"] == 0
        assert s1["total_students"] == 2


def test_reports_sessions_list_sorting_deterministic(teacher_client):
    """Verify completed sessions are sorted deterministically."""
    students, enrollments, sessions, records = _make_phase2_mock_data()
    # Add a third session with same ended_at and started_at
    sessions.append({
        "id": "sess-sub-z-999",
        "subject_id": "subj-1",
        "classroom_id": "class-101",
        "teacher_id": "teacher-smith",
        "started_at": "2026-07-12T11:00:00Z",
        "ended_at": "2026-07-12T12:00:00Z",
        "status": "completed",
        "classroom": {"name": "Lab 201"},
        "teacher": {"full_name": "Dr. Smith"},
        "subject": {"name": "Data Structures"},
    })

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/sessions")
        assert response.status_code == 200
        data = response.json()

        items = data["items"]
        assert items[0]["session_id"] == "sess-sub-z-999"
        assert items[1]["session_id"] == "sess-cust-2"
        assert items[2]["session_id"] == "sess-sub-1"


def test_reports_sessions_list_search(teacher_client):
    """Verify search filter works case-insensitively on subject, classroom, and teacher names."""
    students, enrollments, sessions, records = _make_phase2_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(USER_REPO_MOCK, return_value=students),
        patch(SESS_REPO_MOCK, return_value=sessions),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        # Search by subject name "structures"
        response = teacher_client.get("/api/v1/reports/sessions", params={"search": "structures"})
        assert response.status_code == 200
        data = response.json()
        assert len(data["items"]) == 1
        assert data["items"][0]["session_id"] == "sess-sub-1"

        # Search by classroom name "seminar"
        response = teacher_client.get("/api/v1/reports/sessions", params={"search": "seminar"})
        assert response.status_code == 200
        data = response.json()
        assert len(data["items"]) == 1
        assert data["items"][0]["session_id"] == "sess-cust-2"

        # Search by teacher name "jones"
        response = teacher_client.get("/api/v1/reports/sessions", params={"search": "jones"})
        assert response.status_code == 200
        data = response.json()
        assert len(data["items"]) == 1
        assert data["items"][0]["session_id"] == "sess-cust-2"


def test_reports_sessions_list_student_auth_forbidden(student_client):
    """Verify students cannot list session reports."""
    response = student_client.get("/api/v1/reports/sessions")
    assert response.status_code == 403


def test_reports_session_detail_subject_session(teacher_client):
    """Verify session details for a subject-based session match calculations and structure."""
    students, enrollments, sessions, records = _make_phase2_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(SESS_REL_MOCK, return_value=sessions[0]),
        patch(USER_REPO_MOCK, return_value=students),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/session/sess-sub-1")
        assert response.status_code == 200
        data = response.json()

        # Verify Session metadata
        session = data["session"]
        assert session["session_id"] == "sess-sub-1"
        assert session["subject_name"] == "Data Structures"
        assert session["classroom_name"] == "Lab 201"
        assert session["teacher_name"] == "Dr. Smith"

        # Verify Summary counts
        summary = data["summary"]
        assert summary["present_count"] == 1
        assert summary["late_count"] == 0
        assert summary["absent_count"] == 1
        assert summary["manual_count"] == 0
        assert summary["total_students"] == 2

        # Verify Student roster details (Student 1 present, Student 2 absent)
        students_list = data["students"]
        assert len(students_list) == 2

        s1 = students_list[0]
        assert s1["student_id"] == "student-1"
        assert s1["roll_number"] == "22BCS001"
        assert s1["attendance_status"] == "present"
        assert s1["verification_method"] == "automatic"
        assert s1["marked_at"] == "2026-07-12T09:05:00Z"

        s2 = students_list[1]
        assert s2["student_id"] == "student-2"
        assert s2["roll_number"] == "22BCS002"
        assert s2["attendance_status"] == "absent"
        assert s2["verification_method"] is None
        assert s2["marked_at"] is None


def test_reports_session_detail_custom_session(teacher_client):
    """Verify session details for a custom session (no virtual absences)."""
    students, enrollments, sessions, records = _make_phase2_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(SESS_REL_MOCK, return_value=sessions[1]),
        patch(USER_REPO_MOCK, return_value=students),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get("/api/v1/reports/session/sess-cust-2")
        assert response.status_code == 200
        data = response.json()

        # Verify Session metadata
        session = data["session"]
        assert session["session_id"] == "sess-cust-2"
        assert session["subject_name"] == "Unknown Subject"

        # Verify Summary counts
        summary = data["summary"]
        assert summary["present_count"] == 1
        assert summary["late_count"] == 1
        assert summary["absent_count"] == 0
        assert summary["manual_count"] == 1
        assert summary["total_students"] == 2

        # Verify Student roster details
        students_list = data["students"]
        assert len(students_list) == 2

        s1 = students_list[0]
        assert s1["student_id"] == "student-1"
        assert s1["attendance_status"] == "present"
        assert s1["verification_method"] == "manual"

        s2 = students_list[1]
        assert s2["student_id"] == "student-2"
        assert s2["attendance_status"] == "late"
        assert s2["verification_method"] == "automatic"


def test_reports_session_detail_student_auth_forbidden(student_client):
    """Verify students cannot get session details."""
    response = student_client.get("/api/v1/reports/session/sess-sub-1")
    assert response.status_code == 403


def test_reports_session_detail_not_found(teacher_client):
    """Verify unknown or non-completed session IDs return HTTP 404."""
    with patch(SESS_REL_MOCK, return_value=None):
        response = teacher_client.get("/api/v1/reports/session/unknown-session-id")
        assert response.status_code == 404
        assert response.json()["error"]["code"] == "NOT_FOUND"


def test_reports_session_detail_manual_override_consistency(teacher_client):
    """Verify that updating records updates report results immediately on next request."""
    students, enrollments, sessions, records = _make_phase2_mock_data()

    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(SESS_REL_MOCK, return_value=sessions[0]),
        patch(USER_REPO_MOCK, return_value=students),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        # Initial check
        resp = teacher_client.get("/api/v1/reports/session/sess-sub-1")
        assert resp.json()["summary"]["present_count"] == 1
        assert resp.json()["summary"]["absent_count"] == 1
        assert resp.json()["summary"]["manual_count"] == 0

        # Override records in memory
        records.append({
            "session_id": "sess-sub-1",
            "student_id": "student-2",
            "status": "present",
            "verification_method": "manual",
            "marked_at": "2026-07-12T10:10:00Z",
        })

        # Request again (should reflect new record immediately)
        resp2 = teacher_client.get("/api/v1/reports/session/sess-sub-1")
        assert resp2.json()["summary"]["present_count"] == 2
        assert resp2.json()["summary"]["absent_count"] == 0
        assert resp2.json()["summary"]["manual_count"] == 1

        # Verify student-2 row is updated
        student_rows = resp2.json()["students"]
        s2 = next(x for x in student_rows if x["student_id"] == "student-2")
        assert s2["attendance_status"] == "present"
        assert s2["verification_method"] == "manual"
        assert s2["marked_at"] == "2026-07-12T10:10:00Z"


def test_reports_export_succeeds(teacher_client):
    """Verify that export endpoint succeeds for a teacher and generates styled workbook."""
    students, enrollments, sessions, records = _make_phase2_mock_data()


    def mock_att(ids):
        return [r for r in records if r["session_id"] in ids]

    with (
        patch(SESS_REL_MOCK, return_value=sessions[0]),
        patch(USER_REPO_MOCK, return_value=students),
        patch(ENRL_REPO_MOCK, return_value=enrollments),
        patch(ATTD_REPO_MOCK, side_effect=mock_att),
    ):
        response = teacher_client.get(
            "/api/v1/reports/session/sess-sub-1/export", params={"format": "xlsx"}
        )
        assert response.status_code == 200
        content_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert response.headers["content-type"] == content_type

        # Verify dynamic filename
        content_disp = response.headers["content-disposition"]
        assert 'filename="Data Structures_2026-07-12.xlsx"' in content_disp


        # Load workbook from bytes
        wb = openpyxl.load_workbook(io.BytesIO(response.content))

        # Verify exactly one worksheet named "Session Report"
        assert len(wb.sheetnames) == 1
        assert wb.sheetnames[0] == "Session Report"
        ws = wb["Session Report"]

        # Verify workbook properties
        assert wb.properties.title == "Session Report"
        assert wb.properties.creator == "Smart Campus Attendance System"
        assert wb.properties.subject == "Attendance Report"
        assert wb.properties.description == "Generated attendance report for completed session."

        # Verify metadata values
        assert ws["B1"].value == "Data Structures"
        assert ws["B2"].value == "Dr. Smith"
        assert ws["B3"].value == "Lab 201"

        # Verify metadata time format
        assert ws["B4"].value == "12-Jul-2026 09:00 AM"
        assert ws["B5"].value == "12-Jul-2026 10:00 AM"

        # Verify summary stats
        assert ws["B8"].value == 1
        assert ws["B9"].value == 0
        assert ws["B10"].value == 1
        assert ws["B11"].value == 0
        assert ws["B12"].value == 2

        # Verify headers in row 14
        assert ws["A14"].value == "Roll Number"
        assert ws["B14"].value == "Student Name"
        assert ws["C14"].value == "Attendance Status"
        assert ws["D14"].value == "Verification Method"
        assert ws["E14"].value == "Time Marked"

        # Verify font styles
        assert ws["A14"].font.bold is True
        assert ws["A1"].font.bold is True
        assert ws["A7"].font.bold is True

        # Verify student records count and data rows (row 15 & 16)
        assert ws["A15"].value == "22BCS001"
        assert ws["B15"].value == "Rahul Sharma"
        assert ws["C15"].value == "present"
        assert ws["D15"].value == "automatic"
        assert ws["E15"].value == "12-Jul-2026 09:05 AM"

        assert ws["A16"].value == "22BCS002"
        assert ws["B16"].value == "Priya Patel"
        assert ws["C16"].value == "absent"
        assert ws["D16"].value is None
        assert ws["E16"].value is None


        # Verify scroll freeze pane is applied below headers row
        assert ws.freeze_panes == "A15"


def test_reports_export_student_auth_forbidden(student_client):
    """Verify students are blocked from exporting reports."""
    response = student_client.get("/api/v1/reports/session/sess-sub-1/export")
    assert response.status_code == 403


def test_reports_export_unknown_session_not_found(teacher_client):
    """Verify export requests for unknown sessions return HTTP 404."""
    with patch(SESS_REL_MOCK, return_value=None):
        response = teacher_client.get("/api/v1/reports/session/unknown-session-id/export")
        assert response.status_code == 404


def test_reports_export_unsupported_format(teacher_client):
    """Verify export formats other than xlsx return HTTP 400."""
    response = teacher_client.get(
        "/api/v1/reports/session/sess-sub-1/export", params={"format": "csv"}
    )
    assert response.status_code == 400
    assert "Unsupported export format" in response.json()["error"]["message"]

